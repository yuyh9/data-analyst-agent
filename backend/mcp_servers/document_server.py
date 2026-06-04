import json
import re
import os
from pathlib import Path
from mcp.server import Server
from mcp.server.stdio import stdio_server
from mcp.types import Tool, TextContent
from datetime import datetime

app = Server("document-server")

OUTPUT_DIR = Path("data/reports")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def clean_markdown(text: str) -> str:
    """Strip markdown syntax for plain text contexts."""
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    return text


def md_to_reportlab(text: str) -> str:
    """Convert inline markdown bold to ReportLab XML bold tags."""
    safe = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    safe = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', safe)
    return safe


def generate_docx(title: str, content: str, filename: str) -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # title
    title_para = doc.add_heading(clean_markdown(title), 0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # timestamp
    ts = doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    ts.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ts.runs[0].font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    ts.runs[0].font.size = Pt(10)

    doc.add_paragraph()

    for line in content.split('\n'):
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith('### '):
            doc.add_heading(clean_markdown(line[4:]), level=3)
        elif line.startswith('## '):
            doc.add_heading(clean_markdown(line[3:]), level=2)
        elif line.startswith('# '):
            doc.add_heading(clean_markdown(line[2:]), level=1)
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(clean_markdown(line[2:]), style='List Bullet')
        elif re.match(r'^\d+\.\s', line):
            doc.add_paragraph(clean_markdown(re.sub(r'^\d+\.\s', '', line)), style='List Number')
        else:
            p = doc.add_paragraph()
            # handle inline bold
            parts = re.split(r'\*\*(.*?)\*\*', line)
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.bold = True

    filepath = OUTPUT_DIR / f"{filename}.docx"
    doc.save(str(filepath))
    return str(filepath)


def generate_pdf(title: str, content: str, filename: str) -> str:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    filepath = OUTPUT_DIR / f"{filename}.pdf"
    doc = SimpleDocTemplate(
        str(filepath), pagesize=A4,
        rightMargin=2*cm, leftMargin=2*cm,
        topMargin=2*cm, bottomMargin=2*cm
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('CT', parent=styles['Title'],
        fontSize=20, spaceAfter=8, alignment=TA_CENTER,
        textColor=colors.HexColor('#111827'))
    ts_style = ParagraphStyle('TS', parent=styles['Normal'],
        fontSize=10, spaceAfter=20, alignment=TA_CENTER,
        textColor=colors.HexColor('#9CA3AF'))
    h2_style = ParagraphStyle('H2', parent=styles['Heading2'],
        fontSize=14, spaceBefore=12, spaceAfter=6,
        textColor=colors.HexColor('#1F2937'))
    h3_style = ParagraphStyle('H3', parent=styles['Heading3'],
        fontSize=12, spaceBefore=8, spaceAfter=4,
        textColor=colors.HexColor('#374151'))
    body_style = ParagraphStyle('Body', parent=styles['Normal'],
        fontSize=11, spaceAfter=6, leading=16,
        textColor=colors.HexColor('#374151'))
    bullet_style = ParagraphStyle('Bullet', parent=styles['Normal'],
        fontSize=11, spaceAfter=4, leftIndent=20,
        bulletIndent=10, leading=16,
        textColor=colors.HexColor('#374151'))

    story = []
    story.append(Paragraph(clean_markdown(title), title_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ts_style))

    for line in content.split('\n'):
        line = line.strip()
        if not line:
            story.append(Spacer(1, 6))
            continue
        if line.startswith('## '):
            story.append(Paragraph(clean_markdown(line[3:]), h2_style))
        elif line.startswith('### '):
            story.append(Paragraph(clean_markdown(line[4:]), h3_style))
        elif line.startswith('- ') or line.startswith('* '):
            story.append(Paragraph(f"• {md_to_reportlab(line[2:])}", bullet_style))
        elif re.match(r'^\d+\.\s', line):
            num = re.match(r'^(\d+)\.\s(.*)', line)
            if num:
                story.append(Paragraph(f"{num.group(1)}. {md_to_reportlab(num.group(2))}", body_style))
        else:
            story.append(Paragraph(md_to_reportlab(line), body_style))

    doc.build(story)
    return str(filepath)


@app.list_tools()
async def list_tools() -> list[Tool]:
    return [
        Tool(
            name="export_report",
            description="Export the current analysis as a Word (.docx) or PDF report. Use when user asks to export, save, download, or generate a report.",
            inputSchema={
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Report title"},
                    "content": {"type": "string", "description": "Full report content in markdown format"},
                    "format": {"type": "string", "enum": ["docx", "pdf"], "description": "Output format: docx or pdf"}
                },
                "required": ["title", "content", "format"]
            }
        ),
    ]


@app.call_tool()
async def call_tool(name: str, arguments: dict) -> list[TextContent]:
    try:
        if name == "export_report":
            title = arguments["title"]
            content = arguments["content"]
            fmt = arguments.get("format", "pdf")
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"report_{timestamp}"

            if fmt == "docx":
                filepath = generate_docx(title, content, filename)
            else:
                filepath = generate_pdf(title, content, filename)

            return [TextContent(type="text", text=json.dumps({
                "success": True,
                "filepath": filepath,
                "filename": f"{filename}.{fmt}",
                "format": fmt,
                "title": title
            }))]

        return [TextContent(type="text", text=json.dumps({"error": f"Unknown tool: {name}"}))]

    except Exception as e:
        return [TextContent(type="text", text=json.dumps({"error": str(e)}))]


async def main():
    async with stdio_server() as (read_stream, write_stream):
        await app.run(read_stream, write_stream, app.create_initialization_options())


if __name__ == "__main__":
    import asyncio
    asyncio.run(main())